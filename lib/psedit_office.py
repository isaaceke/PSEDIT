"""
psedit_office — COM-first Office manipulation for PSEdit v14.

When Office is installed (which it is on this machine, Office 16.0),
edits go through COM so Word/Excel themselves do the saving. Every part
of the original file survives: charts, formulas, pivot tables, conditional
formatting, embedded objects, macros. When Office is absent, unsupported
edits raise — no pretending, no silent corruption.
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional


# ── capability detection (registry-based, no imports, no hangs) ─────────
def _reg_exists(subkey: str) -> bool:
    try:
        import winreg
    except Exception:
        return False
    try:
        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, subkey):
            return True
    except OSError:
        return False


def has_word() -> bool:
    return _reg_exists(r"Word.Application\\CurVer")


def has_excel() -> bool:
    return _reg_exists(r"Excel.Application\\CurVer")


def has_openpyxl() -> bool:
    try:
        import openpyxl  # noqa
        return True
    except Exception:
        return False


def has_docx() -> bool:
    try:
        import docx  # noqa
        return True
    except Exception:
        return False


def describe_capabilities() -> str:
    lines = [
        "psedit_office capabilities:",
        f"  word (COM)        {'yes' if has_word() else 'no'}",
        f"  excel (COM)       {'yes' if has_excel() else 'no'}",
        f"  openpyxl          {'yes' if has_openpyxl() else 'no'}",
        f"  python-docx       {'yes' if has_docx() else 'no'}",
    ]
    return chr(10).join(lines)


class OfficeError(Exception): pass
class MissingDependency(OfficeError): pass
class EditFailed(OfficeError): pass


# ── reading ─────────────────────────────────────────────────────────────
def read_xlsx(path: str) -> dict:
    if not has_openpyxl():
        raise MissingDependency("openpyxl not installed")
    from openpyxl import load_workbook
    p = Path(path)
    if not p.exists():
        raise OfficeError(f"file not found: {p}")
    wb = load_workbook(str(p), data_only=False, keep_links=False, keep_vba=False)
    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        merged = [str(r) for r in ws.merged_cells.ranges]
        rows = []
        for row in ws.iter_rows():
            r = []
            for c in row:
                try:
                    is_d = bool(c.is_date)
                except Exception:
                    is_d = False
                try:
                    fmt = c.number_format
                except Exception:
                    fmt = ""
                r.append({"addr": c.coordinate, "value": c.value,
                          "type": type(c.value).__name__,
                          "is_date": is_d, "fmt": fmt})
            rows.append(r)
        sheets.append({"name": name, "max_row": ws.max_row,
                       "max_column": ws.max_column,
                       "merged_cells": merged, "rows": rows})
    wb.close()
    return {"kind": "xlsx", "path": str(p), "sheets": sheets}


def read_docx(path: str) -> dict:
    if not has_docx():
        raise MissingDependency("python-docx not installed")
    from docx import Document
    p = Path(path)
    if not p.exists():
        raise OfficeError(f"file not found: {p}")
    doc = Document(str(p))
    paras = []
    for i, para in enumerate(doc.paragraphs):
        runs = [{"text": r.text, "bold": bool(r.bold),
                 "italic": bool(r.italic), "underline": bool(r.underline)}
                for r in para.runs]
        paras.append({"index": i,
                      "style": para.style.name if para.style else None,
                      "text": para.text, "runs": runs})
    tables = []
    for ti, table in enumerate(doc.tables):
        trows = [[cell.text for cell in row.cells] for row in table.rows]
        tables.append({"index": ti, "rows": trows})
    return {"kind": "docx", "path": str(p),
            "paragraphs": paras, "tables": tables}


# ── markdown formatters ─────────────────────────────────────────────────
def xlsx_to_markdown(data: dict, sheet_name: Optional[str] = None,
                     max_rows: int = 800, max_cols: int = 200) -> str:
    out = []
    for sh in data["sheets"]:
        if sheet_name and sh["name"] != sheet_name:
            continue
        out.append(f"> sheet: {sh['name']}")
        out.append(f"> range: {sh['max_row']} x {sh['max_column']}")
        if sh["merged_cells"]:
            out.append(f"> merged: {', '.join(sh['merged_cells'][:40])}")
        out.append("")
        for row in sh["rows"][:max_rows]:
            cells = []
            for c in row[:max_cols]:
                v = c["value"]
                cells.append("" if v is None else str(v).replace("|", "\\\\|"))
            out.append("| " + " | ".join(cells) + " |")
        out.append("")
    return chr(10).join(out)


def docx_to_markdown(data: dict) -> str:
    out = [f"# {Path(data['path']).name}", ""]
    for para in data["paragraphs"]:
        if para["text"].strip():
            out.append(para["text"])
            out.append("")
    for t in data["tables"]:
        out.append(f"## Table {t['index'] + 1}")
        out.append("")
        for row in t["rows"]:
            out.append("| " + " | ".join(row) + " |")
        out.append("")
    return chr(10).join(out)


# ── COM plumbing ────────────────────────────────────────────────────────
def _com_ready() -> bool:
    try:
        import win32com.client  # noqa
        import pythoncom        # noqa
        return True
    except Exception:
        return False


def _com_init():
    import pythoncom
    pythoncom.CoInitialize()


def _com_uninit():
    try:
        import pythoncom
        pythoncom.CoUninitialize()
    except Exception:
        pass


# ── Excel via COM ───────────────────────────────────────────────────────
def edit_xlsx_com(path: str, replacements: list) -> dict:
    """Open xlsx in Excel, apply cell edits, save. Preserves everything.

    replacements: list of {sheet, cell, value}.
    """
    if not has_excel() or not _com_ready():
        raise MissingDependency("Excel COM unavailable")
    import win32com.client
    p = Path(path).resolve()
    if not p.exists():
        raise OfficeError(f"file not found: {p}")

    applied, errors = 0, []
    excel = wb = None
    _com_init()
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.EnableEvents = False
        try:
            wb = excel.Workbooks.Open(str(p), UpdateLinks=0)
        except Exception as e:
            raise EditFailed(f"Excel refused to open: {e}")
        for op in replacements:
            try:
                ws = wb.Worksheets(op["sheet"])
                ws.Range(op["cell"]).Value = op.get("value")
                applied += 1
            except Exception as e:
                errors.append(f"{op.get('cell')}: {e}")
        if applied:
            wb.Save()
        wb.Close(SaveChanges=False); wb = None
    finally:
        try:
            if wb is not None: wb.Close(SaveChanges=False)
        except Exception: pass
        try:
            if excel is not None: excel.Quit()
        except Exception: pass
        _com_uninit()
    return {"ok": not errors, "applied": applied, "errors": errors}


# ── Word via COM ────────────────────────────────────────────────────────
def edit_docx_com(path: str, replacements: list) -> dict:
    """Open docx in Word, apply find/replace, save. Word's own find handles
    matches that span multiple <w:r> elements — the run-splitting problem
    that python-docx cannot solve.

    replacements: list of {find, replace, match_case}.
    """
    if not has_word() or not _com_ready():
        raise MissingDependency("Word COM unavailable")
    import win32com.client
    p = Path(path).resolve()
    if not p.exists():
        raise OfficeError(f"file not found: {p}")

    applied, errors = 0, []
    word = doc = None
    _com_init()
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        try:
            doc = word.Documents.Open(str(p), ReadOnly=False)
        except Exception as e:
            raise EditFailed(f"Word refused to open: {e}")
        for op in replacements:
            try:
                f = doc.Content.Find
                f.ClearFormatting()
                f.Replacement.ClearFormatting()
                f.Text = op["find"]
                f.Replacement.Text = op.get("replace", "")
                f.Forward = True
                f.Wrap = 1
                f.MatchCase = bool(op.get("match_case", False))
                f.MatchWholeWord = False
                f.MatchWildcards = False
                if f.Execute(Replace=2):
                    applied += 1
                else:
                    errors.append(f"no match: {op['find'][:60]!r}")
            except Exception as e:
                errors.append(f"{op}: {e}")
        if applied:
            doc.Save()
        doc.Close(SaveChanges=False); doc = None
    finally:
        try:
            if doc is not None: doc.Close(SaveChanges=False)
        except Exception: pass
        try:
            if word is not None: word.Quit()
        except Exception: pass
        _com_uninit()
    return {"ok": not errors, "applied": applied, "errors": errors}


# ── validation via COM ──────────────────────────────────────────────────
def validate_office(path: str) -> dict:
    """The strongest validator available: does the native Office app open it?"""
    p = Path(path).resolve()
    if not p.exists():
        return {"ok": False, "error": "file not found"}
    ext = p.suffix.lower()
    if ext in (".xlsx", ".xlsm") and has_excel():
        return _validate_com(p, "Excel.Application", "Workbooks")
    if ext == ".docx" and has_word():
        return _validate_com(p, "Word.Application", "Documents")
    return {"ok": False, "error": f"no validator for {ext}"}


def _validate_com(p: Path, app: str, coll: str) -> dict:
    import win32com.client
    _com_init()
    instance = item = None
    try:
        instance = win32com.client.DispatchEx(app)
        try:
            instance.Visible = False
        except Exception:
            pass
        try:
            instance.DisplayAlerts = 0
        except Exception:
            pass
        item = getattr(instance, coll).Open(str(p), ReadOnly=True)
        item.Close(SaveChanges=False); item = None
        return {"ok": True}
    except Exception as e:
        return {"ok": False, "error": str(e)}
    finally:
        try:
            if item is not None: item.Close(SaveChanges=False)
        except Exception: pass
        try:
            if instance is not None: instance.Quit()
        except Exception: pass
        _com_uninit()


# ── high-level dispatch ─────────────────────────────────────────────────
def read_any(path: str, mode: str = "markdown",
             sheet: Optional[str] = None) -> str:
    import json
    ext = Path(path).suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        data = read_xlsx(path)
        if mode == "markdown":
            return xlsx_to_markdown(data, sheet_name=sheet)
        return json.dumps(data, default=str, ensure_ascii=False, indent=2)
    if ext == ".docx":
        data = read_docx(path)
        if mode == "markdown":
            return docx_to_markdown(data)
        return json.dumps(data, ensure_ascii=False, indent=2)
    raise OfficeError(f"unsupported extension: {ext}")


# ── entry point for standalone testing ──────────────────────────────────
if __name__ == "__main__":
    import sys
    print(describe_capabilities())
    if len(sys.argv) > 1:
        target = sys.argv[1]
        mode = sys.argv[2] if len(sys.argv) > 2 else "markdown"
        sheet = sys.argv[3] if len(sys.argv) > 3 else None
        print(read_any(target, mode=mode, sheet=sheet))